import os #I used os for the path

from customtkinter import * #I used customtkinter because of the UI of this is more appealing than the original
from PIL import Image #I used pillow because the terminal recommended me to due to high pixels of the background and the image 

from openpyxl import Workbook, load_workbook #pyxl imported for storing data from python to excel
from tkinter import messagebox #I also import the normal tkinter for the messagebox

# ----------------------OPENPYXL----------------------------- #

# i used a variable in excel file name so i would not type the file name everytime
filename = "student_scores.xlsx" 

def excel_open(): #This function is the one who will make the excel file in this path and automaticaly type the Name, score and results
    if not os.path.exists(filename):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Score", "Result"])
        wb.save(filename)

def calculate_result(score): #This function will calculate if the user is pass or not 
    return "Pass" if score >= 75 else "Fail"
    

def add_student_score(name, score): #This function will add student info and shows a messagebox if something is wrong like invalid output and success!
    try:
        score = int(score)
    except ValueError:
        messagebox.showerror("Invalid Input", "Score must be a number.")
        return

    result = calculate_result(score)
    wb = load_workbook(filename)
    ws = wb.active
    ws.append([name, score, result])
    wb.save(filename)
    messagebox.showinfo("Success", f"Saved: {name} - {score} ({result})")

def view_all_records(): #This function is when the user click the button it while show the messagebox that have all infos written in the excel file
    if not os.path.exists(filename):
        messagebox.showinfo("No Data", "No records found.")
        return

    wb = load_workbook(filename)
    ws = wb.active
    records = "\n".join([f"{row[0]} | {row[1]} | {row[2]}" for row in ws.iter_rows(min_row=2, values_only=True)])
    messagebox.showinfo("All Records", records if records else "No data yet.")

excel_open() #calling function so if the user run the program it will immediately create a excel file to store the user's info

# ----------------------TKINTER----------------------------- #

mainWindow = CTk()
mainWindow.geometry("400x500")
mainWindow.title("Student's Score Tracker")
set_appearance_mode("light") #this will change the theme of the windows
mainWindow.resizable(False,False) # i used this so the user cannot change the window size

# Background image to make the window more appealing
# i edited the backgound in canva so the backgound won't be plain
bg_image = CTkImage(Image.open("Media Files\\Login Page.png"), size=(400, 500) ) #i sized the image same as the window so it would be not look like cropped or stretched
bg_label = CTkLabel(master=mainWindow, image=bg_image, text="")  #I use label so i can import the image to the window
bg_label.place(x=0, y=0, relwidth=1, relheight=1)

# Same as this one too
jabi_image = CTkImage(light_image=Image.open("Media Files\\miniPicture.png"), size=(75, 75)) #This image is my profile in github just to proof that i made it :)
jabi_label = CTkLabel(master=mainWindow, image=jabi_image, text="")
jabi_label.place(x=160, y=61)  

# The 3 text in the upper part
text1 = CTkLabel(mainWindow, text="Student's",fg_color="white",font=("Helvetica",28, "bold"))
text1.place(x=135,y=140)
text2 = CTkLabel(mainWindow, text="Score Tracker",fg_color="white",font=("Helvetica",28, "bold"))
text2.place(x=100,y=170)
text3 =CTkLabel(mainWindow, text="Made by Jabiii",fg_color="white",font=("Helvetica",12), text_color="grey")
text3.place(x=160,y=200) 

#Here i prioritize the design so it's not boring to look at. I also got the fg color in google (hex) so i can place more colors and make it more appealing.
studentName =CTkLabel(mainWindow, text="Student Name", fg_color="white",font=("Arial",12), text_color="grey")
studentName.place(x=105,y=235)
nameEntry = CTkEntry(master=mainWindow, placeholder_text='Surname, Name, Middle In.' ,height=35, width=180, fg_color="#c4c3be", border_color="#c4c3be")
nameEntry.place(x= 105, y= 260)   

#Same explanation as the upper one but i used the lighter color in entry for appealing
score =CTkLabel(mainWindow, text="Quiz/Test Score",fg_color="white",font=("Arial",12), text_color="grey")
score.place(x=105,y=295)
scoreEntry = CTkEntry(master=mainWindow, placeholder_text="Enter your score",height=35, width=180, fg_color="lightgray", border_color="lightgray")
scoreEntry.place(x= 105, y= 320)

#This function save the infos in excel when the user click the save button after typing infos and if the user didn't type anything the messagebox will occur
def save_data():
    name = nameEntry.get()
    score = scoreEntry.get()
    
    if name.strip() == "" or score.strip() == "":
        messagebox.showwarning("Missing Info", "Please enter both name and score.")
        return
    
    add_student_score(name, score)
    nameEntry.delete(0, END)
    scoreEntry.delete(0, END)

#This Funciton will call the another function that will show all the records. I made it so when i read it i will immediately understand what command it is
def show_data():
    view_all_records()


saveBtn = CTkButton(master=mainWindow, text="Save", fg_color="black", height=35, width=80, command=save_data)
saveBtn.place(x=105, y=365)

viewBtn = CTkButton(master=mainWindow, text="View Data", fg_color="black", height=35, width=80, command=show_data)
viewBtn.place(x=205, y=365)

mainWindow.mainloop() #this will call the window to show it to the monitor