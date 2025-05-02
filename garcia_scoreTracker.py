import os
from customtkinter import *
from PIL import Image, ImageTk

from openpyxl import Workbook, load_workbook
from tkinter import messagebox 

# ----------------------OPENPYXL----------------------------- #

filename = "student_scores.xlsx"

def initialize_file():
    if not os.path.exists(filename):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Score", "Result"])
        wb.save(filename)

def calculate_result(score):
    return "Pass" if score >= 75 else "Fail"
    

def add_student_score(name, score):
    try:
        score = int(score)
    except ValueError:
        messagebox.showerror("Invalid Input", "Score must be a number.")
        return

    result = calculate_result(score)
    wb = load_workbook(filename)
    ws = wb.active
    ws.append([name, score, result])
    wb.save(filename)
    messagebox.showinfo("Success", f"Saved: {name} - {score} ({result})")

def view_all_records():
    if not os.path.exists(filename):
        messagebox.showinfo("No Data", "No records found.")
        return

    wb = load_workbook(filename)
    ws = wb.active
    records = "\n".join([f"{row[0]} | {row[1]} | {row[2]}" for row in ws.iter_rows(min_row=2, values_only=True)])
    messagebox.showinfo("All Records", records if records else "No data yet.")

initialize_file()

# ----------------------TKINTER----------------------------- #

mainWindow = CTk()
mainWindow.geometry("400x500")
mainWindow.title("Student's Score Tracker")
set_appearance_mode("light")
mainWindow.resizable(False,False)

bg_image = CTkImage(Image.open("Media Files\\Login Page.png"), size=(400, 500) )
bg_label = CTkLabel(master=mainWindow, image=bg_image, text="")  
bg_label.place(x=0, y=0, relwidth=1, relheight=1)


jabi_image = CTkImage(light_image=Image.open("Media Files\\miniPicture.png"), size=(75, 75))
jabi_label = CTkLabel(master=mainWindow, image=jabi_image, text="")
jabi_label.place(x=160, y=61)  


text1 = CTkLabel(mainWindow, text="Student's",fg_color="white",font=("Helvetica",28, "bold"))
text1.place(x=135,y=140)
text2 = CTkLabel(mainWindow, text="Score Tracker",fg_color="white",font=("Helvetica",28, "bold"))
text2.place(x=100,y=170)
text3 =CTkLabel(mainWindow, text="Made by Jabiii",fg_color="white",font=("Helvetica",12), text_color="grey")
text3.place(x=160,y=200)


studentName =CTkLabel(mainWindow, text="Student Name", fg_color="white",font=("Arial",12), text_color="grey")
studentName.place(x=105,y=235)
nameEntry = CTkEntry(master=mainWindow, placeholder_text='Surname, Name, Middle In.' ,height=35, width=180, fg_color="#c4c3be", border_color="#c4c3be")
nameEntry.place(x= 105, y= 260)   

score =CTkLabel(mainWindow, text="Quiz/Test Score",fg_color="white",font=("Arial",12), text_color="grey")
score.place(x=105,y=295)
scoreEntry = CTkEntry(master=mainWindow, placeholder_text="Enter your score",height=35, width=180, fg_color="lightgray", border_color="lightgray")
scoreEntry.place(x= 105, y= 320)

def save_data():
    name = nameEntry.get()
    score = scoreEntry.get()
    
    if name.strip() == "" or score.strip() == "":
        messagebox.showwarning("Missing Info", "Please enter both name and score.")
        return
    
    add_student_score(name, score)
    nameEntry.delete(0, END)
    scoreEntry.delete(0, END)

def show_data():
    view_all_records()


saveBtn = CTkButton(master=mainWindow, text="Save", fg_color="black", height=35, width=80, command=save_data)
saveBtn.place(x=105, y=365)

viewBtn = CTkButton(master=mainWindow, text="View Data", fg_color="black", height=35, width=80, command=show_data)
viewBtn.place(x=205, y=365)



mainWindow.mainloop()